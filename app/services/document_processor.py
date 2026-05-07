import fitz  # PyMuPDF
from paddleocr import PaddleOCR
from PIL import Image
import numpy as np
import os
import subprocess
import tempfile
from pathlib import Path
from docx import Document as DocxDocument
from enum import Enum
from dataclasses import dataclass
from typing import List, Dict, Optional, Union
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentType(Enum):
    NATIVE_PDF = "native_pdf"
    SCANNED_PDF = "scanned_pdf"
    IMAGE = "image"
    WORD = "word"
    UNKNOWN = "unknown"

@dataclass
class ExtractedPage:
    page_number: int
    text: str
    tables: List[Dict]
    confidence: float
    bounding_boxes: List[Dict]  # For audit trail
    ocr_required: bool

@dataclass
class ProcessedDocument:
    document_id: str
    original_filename: str
    document_type: DocumentType
    pages: List[ExtractedPage]
    overall_confidence: float
    needs_manual_review: bool
    extraction_metadata: Dict

class DocumentProcessor:
    def __init__(self):
        self.ocr = PaddleOCR(
            use_angle_cls=True,
            lang='en',
            use_gpu=False,  # Set True if GPU available
            show_log=False
        )
        self.confidence_threshold = 0.80
        # Check if LibreOffice is available for Word conversion
        self.libreoffice_available = self._check_libreoffice()

    def _check_libreoffice(self) -> bool:
        """Check if LibreOffice is available for document conversion."""
        try:
            subprocess.run(['soffice', '--version'], 
                          capture_output=True, check=True)
            return True
        except (subprocess.CalledProcessError, FileNotFoundError):
            try:
                subprocess.run(['libreoffice', '--version'], 
                              capture_output=True, check=True)
                return True
            except (subprocess.CalledProcessError, FileNotFoundError):
                logger.warning("LibreOffice not found. Word document processing will be limited to text extraction.")
                return False

    def classify_document(self, file_path: str) -> DocumentType:
        """Determine document type and whether OCR is needed."""
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        suffix = file_path.suffix.lower()
        
        if suffix in ['.png', '.jpg', '.jpeg', '.tiff', '.bmp']:
            return DocumentType.IMAGE
        elif suffix == '.pdf':
            return self._classify_pdf(str(file_path))
        elif suffix in ['.doc', '.docx']:
            return DocumentType.WORD
        else:
            return DocumentType.UNKNOWN

    def _classify_pdf(self, file_path: str) -> DocumentType:
        """Check if PDF has extractable text or is scanned."""
        try:
            doc = fitz.open(file_path)
            text_pages = sum(1 for page in doc if page.get_text().strip())
            # If more than 50% of pages have extractable text, consider it native
            return DocumentType.NATIVE_PDF if text_pages > len(doc) * 0.5 else DocumentType.SCANNED_PDF
        except Exception as e:
            logger.error(f"Error classifying PDF {file_path}: {e}")
            return DocumentType.SCANNED_PDF  # Default to scanned if unsure

    def process(self, file_path: str, document_id: str) -> ProcessedDocument:
        """Main processing pipeline dispatching by document type."""
        doc_type = self.classify_document(file_path)
        
        if doc_type == DocumentType.NATIVE_PDF:
            return self._process_native_pdf(file_path, document_id)
        elif doc_type == DocumentType.SCANNED_PDF:
            return self._process_with_ocr(file_path, document_id, doc_type)
        elif doc_type == DocumentType.IMAGE:
            return self._process_with_ocr(file_path, document_id, doc_type)
        elif doc_type == DocumentType.WORD:
            return self._process_word(file_path, document_id)
        else:
            raise ValueError(f"Unsupported document type: {doc_type}")

    def _process_native_pdf(self, file_path: str, document_id: str) -> ProcessedDocument:
        """Process text-based PDF documents."""
        doc = fitz.open(file_path)
        pages = []
        
        for page_num, page in enumerate(doc):
            text = page.get_text()
            # For native PDF, we assume high confidence for text
            # Tables extraction would go here (simplified for now)
            tables = []  # Placeholder - implement table extraction if needed
            
            pages.append(ExtractedPage(
                page_number=page_num + 1,
                text=text,
                tables=tables,
                confidence=1.0,  # Native PDF text is reliable
                bounding_boxes=[],  # Would need more work for accurate bboxes
                ocr_required=False
            ))
        
        overall_conf = 1.0  # Native PDFs get full confidence
        
        return ProcessedDocument(
            document_id=document_id,
            original_filename=os.path.basename(file_path),
            document_type=DocumentType.NATIVE_PDF,
            pages=pages,
            overall_confidence=overall_conf,
            needs_manual_review=overall_conf < self.confidence_threshold,
            extraction_metadata={
                'processor': 'PyMuPDF',
                'processing_timestamp': str(np.datetime64('now'))
            }
        )

    def _process_with_ocr(self, file_path: str, document_id: str,
                          doc_type: DocumentType) -> ProcessedDocument:
        """Process scanned PDFs or images using OCR."""
        # Convert to images if needed
        if doc_type == DocumentType.SCANNED_PDF:
            images = self._pdf_to_images(file_path)
        else:  # IMAGE
            images = [Image.open(file_path)]
        
        pages = []
        
        for idx, img in enumerate(images):
            # Convert PIL image to numpy array for PaddleOCR
            img_array = np.array(img)
            
            # Run OCR
            result = self.ocr.ocr(img_array, cls=True)
            
            text_blocks = []
            confidences = []
            bboxes = []
            
            if result and result[0]:
                for line in result[0]:
                    bbox, (text, conf) = line
                    text_blocks.append(text)
                    confidences.append(conf)
                    bboxes.append({
                        'text': text,
                        'confidence': conf,
                        'bbox': bbox
                    })
            
            avg_confidence = sum(confidences) / len(confidences) if confidences else 0
            
            # Extract tables (simplified - would need more sophisticated table detection)
            tables = self._extract_tables_from_ocr(bboxes) if bboxes else []
            
            pages.append(ExtractedPage(
                page_number=idx + 1,
                text='\n'.join(text_blocks),
                tables=tables,
                confidence=avg_confidence,
                bounding_boxes=bboxes,
                ocr_required=True
            ))
        
        overall_conf = sum(p.confidence for p in pages) / len(pages) if pages else 0
        
        return ProcessedDocument(
            document_id=document_id,
            original_filename=os.path.basename(file_path),
            document_type=doc_type,
            pages=pages,
            overall_confidence=overall_conf,
            needs_manual_review=overall_conf < self.confidence_threshold,
            extraction_metadata={
                'ocr_engine': 'PaddleOCR',
                'processing_timestamp': str(np.datetime64('now'))
            }
        )

    def _process_word(self, file_path: str, document_id: str) -> ProcessedDocument:
        """Process Word documents by converting to PDF or extracting text."""
        # Try to convert to PDF first if LibreOffice is available
        if self.libreoffice_available:
            try:
                pdf_path = self._convert_word_to_pdf(file_path)
                # Process the resulting PDF
                return self._process_with_ocr(pdf_path, document_id, DocumentType.SCANNED_PDF)
            except Exception as e:
                logger.warning(f"Word to PDF conversion failed: {e}. Falling back to text extraction.")
                # Fall through to text extraction
        
        # Fallback: Extract text directly from Word (no OCR for images)
        return self._extract_word_text(file_path, document_id)

    def _convert_word_to_pdf(self, input_path: str) -> str:
        """Convert Word document to PDF using LibreOffice."""
        with tempfile.TemporaryDirectory() as temp_dir:
            cmd = [
                'soffice' if self.libreoffice_available and self._check_libreoffice_cmd('soffice') else 'libreoffice',
                '--headless',
                '--convert-to',
                'pdf',
                '--outdir',
                temp_dir,
                input_path
            ]
            
            try:
                subprocess.run(cmd, check=True, capture_output=True, timeout=30)
            except subprocess.CalledProcessError as e:
                logger.error(f"LibreOffice conversion failed: {e.stderr.decode()}")
                raise
            except FileNotFoundError:
                # Try the other command
                cmd[0] = 'libreoffice' if cmd[0] == 'soffice' else 'soffice'
                subprocess.run(cmd, check=True, capture_output=True, timeout=30)
            
            # Find the generated PDF
            input_name = Path(input_path).stem
            pdf_path = os.path.join(temp_dir, f"{input_name}.pdf")
            
            if not os.path.exists(pdf_path):
                # Try alternative naming
                for file in os.listdir(temp_dir):
                    if file.endswith('.pdf'):
                        pdf_path = os.path.join(temp_dir, file)
                        break
                else:
                    raise FileNotFoundError("PDF conversion did not produce output file")
            
            # Copy to a persistent temporary location since temp_dir will be cleaned
            persistent_temp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            persistent_temp.close()
            
            import shutil
            shutil.copy2(pdf_path, persistent_temp.name)
            return persistent_temp.name

    def _check_libreoffice_cmd(self, cmd: str) -> bool:
        """Check if a specific LibreOffice command is available."""
        try:
            subprocess.run([cmd, '--version'], 
                          capture_output=True, check=True, timeout=5)
            return True
        except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired):
            return False

    def _extract_word_text(self, file_path: str, document_id: str) -> ProcessedDocument:
        """Extract text from Word document using python-docx (no OCR)."""
        try:
            doc = DocxDocument(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            
            text = '\n'.join(full_text)
            
            # Create a single "page" with the extracted text
            pages = [ExtractedPage(
                page_number=1,
                text=text,
                tables=[],  # Would need more work to extract table structure
                confidence=0.7,  # Medium confidence since we might miss text in images/complex layouts
                bounding_boxes=[],
                ocr_required=False  # We didn't use OCR
            )]
            
            return ProcessedDocument(
                document_id=document_id,
                original_filename=os.path.basename(file_path),
                document_type=DocumentType.WORD,
                pages=pages,
                overall_confidence=0.7,  # Fixed confidence for Word text extraction
                needs_manual_review=0.7 < self.confidence_threshold,  # Will be True if threshold > 0.7
                extraction_metadata={
                    'processor': 'python-docx',
                    'method': 'text_extraction_only',
                    'warning': 'No OCR performed - text in images or complex layouts may be missed',
                    'processing_timestamp': str(np.datetime64('now'))
                }
            )
        except Exception as e:
            logger.error(f"Error extracting text from Word document {file_path}: {e}")
            raise

    def _pdf_to_images(self, pdf_path: str) -> List[Image.Image]:
        """Convert PDF pages to PIL Images."""
        doc = fitz.open(pdf_path)
        images = []
        
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x zoom for better OCR
            img_data = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_data))
            images.append(img)
        
        return images

    def _extract_tables_from_ocr(self, bboxes: List[Dict]) -> List[Dict]:
        """
        Simplified table extraction from OCR bounding boxes.
        In a production system, you'd use more sophisticated table detection.
        """
        # This is a placeholder - implement actual table detection based on bbox alignment
        # For now, return empty list
        return []

# Example usage
if __name__ == "__main__":
    import io  # Needed for _pdf_to_images
    
    processor = DocumentProcessor()
    
    # Test with a sample file - replace with actual path
    test_file = "sample_document.pdf"  # Change to your test file
    
    if os.path.exists(test_file):
        try:
            result = processor.process(test_file, "doc123")
            print(f"Processed {test_file}")
            print(f"Document type: {result.document_type.value}")
            print(f"Pages: {len(result.pages)}")
            print(f"Overall confidence: {result.overall_confidence:.2f}")
            print(f"Needs manual review: {result.needs_manual_review}")
            
            for i, page in enumerate(result.pages):
                print(f"Page {i+1}: {len(page.text)} chars, confidence: {page.confidence:.2f}")
        except Exception as e:
            print(f"Error processing {test_file}: {e}")
    else:
        print(f"Test file {test_file} not found. Please provide a valid document path.")